#!/usr/bin/env python3
"""Génère l'Excel des événements de la Semaine de l'innovation 2026 (3e édition).

Sources : planning PDF « mise à jour 22-10 » (titres, horaires, lieux, séances ouvertes aux entreprises) complété par
les fiches de https://www.u-bordeaux.fr/evenements/ (type, public, adresse, lien).

Usage : python3 data/build_xlsx.py [sortie.xlsx]
Dépendance : openpyxl (uniquement pour ce script, pas pour le widget).
"""
import json
import sys
from calendar import timegm
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

SITE = "https://www.u-bordeaux.fr/evenements/"
PROGRAMME = SITE + "semaine-de-linnovation-2026-3e-edition"

# Les en-têtes sont volontairement sans accent ni espace : ce sont les identifiants
# de colonnes Grist que le widget attend (voir CONFIG.evenements dans widget.js).
COLONNES = [
    ("Jour", 12), ("Debut", 8), ("Fin", 8), ("Titre", 44), ("Sous_titre", 44),
    ("Type", 18), ("Temps_fort", 24), ("Lieu", 38), ("Ville", 16), ("Adresse", 44),
    ("Public", 40), ("Cibles", 30), ("Description", 70), ("Lien", 60),
]

LUN, MAR, MER, JEU, VEN = (date(2026, 11, d) for d in (23, 24, 25, 26, 27))

TEMPS_FORT_GP = "Grand public"
TEMPS_FORT_SE = "Acteurs socio-économiques"

BU_DL = "4 avenue Denis Diderot, 33600 Pessac"
HAUT_CARRE = "43 rue Pierre Noailles, 33400 Talence"
IMS = "351 cours de la Libération, 33400 Talence"
ARTS_METIERS = "Esplanade des Arts et Métiers, 33400 Talence"
BU_ST = "Avenue des Facultés, 33400 Talence"

# Publics visés, d'après les quatre plannings par public (Pédagogie, Recherche, Étudiants, Entreprises ; ce dernier revu d'après
# le planning du 22/10) : une séance est « pour » un public quand elle est surlignée sur son planning. Valeurs de la colonne « Cibles » (liste de choix).
# Cinquième public, « Personnels » : la seule séance qu'aucun planning ne vise (public du site : personnels administratifs).
PEDAGOGIE, RECHERCHE, ETUDIANTS, ENTREPRISES = "Pédagogie", "Recherche", "Étudiants", "Entreprises"
PERSONNELS = "Personnels"
TOUS = (PEDAGOGIE, RECHERCHE, ETUDIANTS, ENTREPRISES)   # les quatre plannings : « tous » ne comprend pas les personnels
ORDRE = TOUS + (PERSONNELS,)

# (jour, début, titre) -> publics. Une séance absente de cette table n'est surlignée sur aucun planning.
CIBLES = {
    (LUN, "09:30", "Le Transmedia Storytelling comme valorisation de projets"): (PEDAGOGIE, RECHERCHE),
    (LUN, "11:30", "Les DéJEUXners de l’innovation"): (PEDAGOGIE, RECHERCHE),
    (LUN, "14:00", "Construisons ensemble un espace pour apprendre autrement"): (PEDAGOGIE,),
    (LUN, "14:00", "Atelier Durable"): (PEDAGOGIE, ETUDIANTS),
    (LUN, "18:00", "Conférence « L’IA au service des patients : quelles innovations en santé ? »"): TOUS,
    (MAR, "09:00", "Entrepreneuriat et innovation au service des transitions agroécologiques"): (ETUDIANTS, ENTREPRISES),
    (MAR, "09:00", "Entractes : regards croisés sur les collaborations entre recherche académique en santé et entreprise"): (RECHERCHE, ENTREPRISES),
    (MAR, "09:00", "Matinale scientifique"): (RECHERCHE, ENTREPRISES),
    (MAR, "11:00", "Les Evolutive Learning Factories, des usines-écoles pour apprendre autrement"): (PEDAGOGIE, ENTREPRISES),
    (MAR, "11:30", "Les DéJEUXners de l’innovation"): (ETUDIANTS,),
    (MAR, "14:00", "De la recherche à l’innovation concrète en entreprise : la valeur ajoutée d’un Centre de Ressources Technologiques comme le CATIE"): (ENTREPRISES,),
    (MAR, "16:00", "Construire autrement : bois de feuillus, et terres allégées"): (RECHERCHE, ETUDIANTS, ENTREPRISES),
    (MAR, "18:30", "Mardis de l’entrepreneuriat"): (RECHERCHE, ETUDIANTS),
    (MER, "09:00", "Faire grandir ses partenariats"): (RECHERCHE, ENTREPRISES),
    (MER, "10:00", "Micro et nanofabrication pour la photonique avancée"): (RECHERCHE, ENTREPRISES),
    (MER, "10:00", "Du laboratoire à l’entreprise, le potentiel de l’Open Source issu de la recherche"): (RECHERCHE, ENTREPRISES),
    (MER, "14:00", "Space Meet-Up Nouvelle-Aquitaine : explorer l’Espace avec nos experts"): (ENTREPRISES,),
    (MER, "17:40", "Les cellules de transfert technologique en pratique"): (RECHERCHE, ENTREPRISES),
    (MER, "18:30", "Salon des entreprises"): (ENTREPRISES,),
    (JEU, "09:00", "Albatros Techno Day 2026"): (RECHERCHE,),
    (JEU, "09:30", "CampusIABordeaux"): (PEDAGOGIE,),
    (JEU, "10:30", "Échoppe des Sciences"): (PEDAGOGIE,),
    (JEU, "11:30", "Les DéJEUXners de l’innovation"): (ETUDIANTS,),
    (JEU, "13:00", "Pépinière numérique FabLabNab"): (PEDAGOGIE, RECHERCHE),
    (JEU, "14:00", "La robotique au service de la fabrication additive métallique"): (PEDAGOGIE, RECHERCHE, ENTREPRISES),
    (JEU, "17:30", "Afterwork docteurs, doctorants et entreprises"): (RECHERCHE, ENTREPRISES),
    (VEN, "09:00", "CampusIABordeaux"): (PEDAGOGIE,),
    (VEN, "09:00", "Des formations sur mesure pour faire évoluer ses pratiques"): (PEDAGOGIE, RECHERCHE),
    (VEN, "11:30", "CampusIABordeaux"): (PEDAGOGIE,),
    # Surligné sur aucun des quatre plannings ; public du site : personnels administratifs.
    (VEN, "11:00", "Serious Game « No wrong door » du Pôle universitaire d’innovation Bordeaux"): (PERSONNELS,),
}


def cibles_de(jour, debut, titre):
    """Publics d'une séance, dans l'ordre Pédagogie, Recherche, Étudiants, Entreprises, Personnels, séparés par des virgules."""
    publics = CIBLES.get((jour, debut, titre), ())
    return ", ".join(c for c in ORDRE if c in publics)


# (jour, début, fin, titre, sous-titre, type, temps fort, lieu, ville, adresse, public, description, fiche)
EVENEMENTS = [
    # ---------------------------------------------------------------- Lundi 23
    (LUN, "09:30", "12:30", "Le Transmedia Storytelling comme valorisation de projets", "",
     "Atelier", "", "BU Droit-Lettres, salle 205", "Pessac", "Salle 205, " + BU_DL,
     "Enseignants-chercheurs, doctorants",
     "Initiation au Transmedia Storytelling : concevoir en groupe une stratégie de récits multiplateformes pour valoriser un projet.",
     "le-transmedia-storytelling-comme-valorisation-de-projets"),
    (LUN, "11:30", "13:00", "Les DéJEUXners de l’innovation", "",
     "Atelier", "", "Haut-Carré, restaurant des personnels", "Talence", HAUT_CARRE,
     "Personnels",
     "Jeux de société détournés et défis de créativité pour découvrir les méthodes de l’innovation de façon ludique, sans prérequis.",
     "les-dejeuxners-de-linnovation"),
    (LUN, "14:00", "16:00", "Construisons ensemble un espace pour apprendre autrement", "",
     "Atelier", "", "BU Sciences et Techniques", "Talence", "Salle « île aux oiseaux », " + BU_ST,
     "Enseignants-chercheurs, accompagnants pédagogiques",
     "Atelier de co-conception pour imaginer des espaces d’apprentissage innovants et inclusifs, à partir de scénarios pédagogiques concrets.",
     "construisons-ensemble-un-espace-pour-apprendre-autrement"),
    (LUN, "14:00", "16:00", "Atelier Durable", "Évaluer pour mieux maîtriser les impacts environnementaux",
     "Atelier", "", "ENSMAC - Bordeaux INP", "Pessac", "16 avenue Pey Berland, 33600 Pessac",
     "Enseignants-chercheurs, chercheurs, personnels, étudiants",
     "Outils ludiques (jeu de société collaboratif, calculateur Ecobalyse) pour évaluer l’empreinte environnementale de produits.",
     "evaluer-pour-mieux-maitriser-les-impacts-environnementaux"),
    (LUN, "18:00", "", "Conférence « L’IA au service des patients : quelles innovations en santé ? »", "",
     "Conférence", TEMPS_FORT_GP, "Amphithéâtre Station Marne, campus Victoire", "Bordeaux", "",
     "Grand public", "", None),
    # ---------------------------------------------------------------- Mardi 24
    (MAR, "09:00", "16:00", "Entrepreneuriat et innovation au service des transitions agroécologiques", "",
     "Conférence", "", "Bordeaux Sciences Agro", "Gradignan", "1 cours du Général de Gaulle, 33170 Gradignan",
     "Étudiants, personnels, acteurs socio-économiques",
     "Journée de présentations, témoignages et échanges sur les dispositifs d’appui à l’entrepreneuriat agricole et agroalimentaire, avec visite de la plateforme FarmLab@Bordeaux.",
     "un-terrain-fertile-pour-lentrepreneuriat-et-linnovation-au-service-des-transitions-agroecologiques"),
    (MAR, "09:00", "12:00", "Entractes : regards croisés sur les collaborations entre recherche académique en santé et entreprise", "",
     "Table ronde", "", "Amphithéâtre BBS, campus Carreire", "Bordeaux", "2 rue du Dr Hoffmann Martinot, 33000 Bordeaux",
     "Enseignants-chercheurs, chercheurs, doctorants, étudiants de master",
     "Des binômes chercheur / industriel partagent leurs retours sur leurs collaborations ; présentation des dispositifs pour monter et financer des projets communs.",
     "entractes-regards-croises-sur-les-collaborations-entre-recherche-academique-en-sante-et-entreprise"),
    (MAR, "09:00", "12:30", "Matinale scientifique", "Performance énergétique et environnementale des bâtiments",
     "Conférence", "", "IMS, amphithéâtre Jean-Paul Dom", "Talence", "Bâtiment A31, " + IMS,
     "Enseignants-chercheurs, acteurs socio-économiques",
     "Matinée organisée par NOBATEK et l’université de Bordeaux (équipe GP2E) sur la performance énergétique et environnementale des bâtiments.",
     "matinale-scientifique-garantie-de-performances-energetiques-et-environnementales-des-batiments"),
    (MAR, "11:00", "12:00", "Les Evolutive Learning Factories, des usines-écoles pour apprendre autrement", "",
     "Visite", "", "Arts et Métiers, halle technologique", "Talence", ARTS_METIERS,
     "Enseignants-chercheurs, ingénieurs, professionnels de la formation",
     "Visite de la Learning Factory des Arts et Métiers : des plateformes pédagogiques proches de l’industrie, du prototypage à la production mécanique.",
     "les-evolutive-learning-factories-des-usines-ecoles-pour-apprendre-autrement"),
    (MAR, "11:30", "13:00", "Les DéJEUXners de l’innovation", "",
     "Atelier", "", "Restaurant universitaire l’Amazone", "Talence", "",
     "Étudiants",
     "Jeux de société détournés et défis de créativité pour découvrir les méthodes de l’innovation de façon ludique, sans prérequis.",
     "les-dejeuxners-de-linnovation"),
    (MAR, "14:00", "16:30", "De la recherche à l’innovation concrète en entreprise : la valeur ajoutée d’un Centre de Ressources Technologiques comme le CATIE", "",
     "Visite", "", "Locaux du CATIE, ENSEIRB-MATMECA", "Talence", "",
     "Acteurs socio-économiques",
     "Visite commentée du CATIE, centre de ressources technologiques du numérique : expertises, équipements et projets menés avec des entreprises.",
     "de-la-recheche-linnovation-concrete-en-entreprise-la-valeur-ajoutee-dun-centre-de-ressources-technologiques-comme-le-catie"),
    (MAR, "16:00", "18:00", "Construire autrement : bois de feuillus, et terres allégées", "",
     "Visite", "", "ENSAP Bordeaux", "Talence", "740 cours de la Libération, 33400 Talence",
     "Architectes, ingénieurs, chercheurs, étudiants, professionnels de la construction et de la filière bois",
     "Visite des démonstrateurs THRIVE et TNEB : construire avec le hêtre et les feuillus locaux, la terre-paille et le chanvre. Départs des visites à 16 h et 17 h.",
     "construire-autrement-avec-le-hetre-la-terre-paille-et-le-chanvre"),
    (MAR, "18:30", "21:30", "Mardis de l’entrepreneuriat", "Les sciences humaines et sociales, moteur d’innovation",
     "Table ronde", "", "BU Droit-Lettres, hall", "Pessac", BU_DL,
     "Étudiants et doctorants en sciences humaines et sociales",
     "Soirée en trois temps : table ronde, atelier collaboratif et pitchs d’étudiants autour de l’innovation en SHS.",
     "mardis-de-lentrepreneuriat-les-sciences-humaines-et-sociales-moteur-dinnovation"),
    # ------------------------------------------------------------- Mercredi 25
    (MER, "09:00", "11:00", "Faire grandir ses partenariats", "Des collaborations structurantes en sciences humaines et sociales",
     "Table ronde", "", "BU Droit-Lettres, salle 205", "Pessac", "Salle 205, " + BU_DL,
     "Chercheurs en sciences humaines et sociales",
     "Transformer une première collaboration avec un acteur socio-économique en partenariat durable : témoignages et échanges avec des experts.",
     "faire-grandir-ses-partenariats-construire-des-collaborations-structurantes-en-sciences-humaines-et-sociales"),
    (MER, "10:00", "17:00", "Micro et nanofabrication pour la photonique avancée", "",
     "Table ronde", "", "Haut-Carré, salle des Conseils", "Talence", "Entrée Agora, " + HAUT_CARRE,
     "Chercheurs, enseignants-chercheurs, industriels",
     "Journée autour des équipements de micro et nanofabrication du site bordelais pour définir une feuille de route technologique en photonique.",
     "micro-et-nanofabrication-pour-la-photonique-avancee"),
    (MER, "10:00", "12:00", "Du laboratoire à l’entreprise, le potentiel de l’Open Source issu de la recherche", "",
     "Table ronde", "", "LaBRI, amphithéâtre, bât. A30", "Talence", "Bâtiment A30, " + IMS,
     "Chercheurs du numérique, personnels de la valorisation, acteurs socio-économiques",
     "Valoriser les logiciels de recherche en Open Source : visibilité, collaborations avec les entreprises, modèles économiques, licences.",
     "du-laboratoire-lentreprise-le-potentiel-de-lopen-source-issu-de-la-recherche"),
    (MER, "14:00", "17:30", "Space Meet-Up Nouvelle-Aquitaine : explorer l’Espace avec nos experts", "",
     "Conférence", "", "IMS, amphithéâtre Jean-Paul Dom", "Talence", "Bâtiment A31, " + IMS,
     "Acteurs socio-économiques du secteur spatial",
     "Découvrir les expertises, laboratoires et plateformes spatiales de l’université de Bordeaux, puis échanger avec les chercheurs (avec Aerospace Valley).",
     "explorer-lespace-avec-nos-experts"),
    (MER, "17:40", "18:30", "Les cellules de transfert technologique en pratique", "",
     "Table ronde", "", "Agora, Domaine Haut-Carré", "Talence", HAUT_CARRE,
     "Enseignants-chercheurs, chercheurs, acteurs socio-économiques",
     "Les coulisses des collaborations laboratoires-entreprises et le rôle des cellules de transfert technologique.",
     "les-cellules-de-transfert-technologique-en-pratique"),
    (MER, "18:30", "20:30", "Salon des entreprises", "Du concept au démonstrateur",
     "Salon", TEMPS_FORT_SE, "Salle Badiane, Domaine du Haut-Carré", "Talence", HAUT_CARRE,
     "Acteurs socio-économiques (start-ups, deeptech, PME, ETI)",
     "Rencontrer les plateformes technologiques de l’université de Bordeaux et découvrir leurs équipements et compétences.",
     "salon-des-entreprises-du-concept-au-demonstrateur"),
    # ---------------------------------------------------------------- Jeudi 26
    (JEU, "09:00", "17:00", "Albatros Techno Day 2026", "",
     "Conférence", "", "IMS, amphithéâtre Jean-Paul Dom", "Talence", "Bâtiment A31, " + IMS,
     "Enseignants-chercheurs, chercheurs, acteurs socio-économiques",
     "Rencontres entre chercheurs du GIS Albatros et industriels pour identifier des sujets communs (micro-projets, études, thèses CIFRE).",
     "albatros-techno-day-2026"),
    (JEU, "09:30", "10:30", "CampusIABordeaux", "Entraînez vos étudiants sur la donnée",
     "Atelier", "", "BU Sciences et Techniques", "Talence", "Salle « île aux oiseaux », " + BU_ST,
     "Enseignants-chercheurs",
     "Présentation de la plateforme CampusIABordeaux et de son usage avec les étudiants : publication, partage et réutilisation des données.",
     "campusiabordeaux-entrainez-vos-etudiants-sur-la-donnee"),
    (JEU, "10:30", "12:15", "Échoppe des Sciences", "Découvrez, questionnez, adoptez !",
     "Atelier", "", "BU Sciences et Techniques", "Talence", "Salle « île aux oiseaux », " + BU_ST,
     "Enseignants-chercheurs",
     "Présentation du dispositif qui accompagne des projets d’acteurs de la société et co-construit des questions de recherche. Accueil café à 10 h 30, atelier à 10 h 45.",
     "echoppe-des-sciences-decouvrez-questionnez-adoptez"),
    (JEU, "11:30", "13:00", "Les DéJEUXners de l’innovation", "",
     "Atelier", "", "Restaurant universitaire n° 2", "Pessac", "",
     "Étudiants",
     "Jeux de société détournés et défis de créativité pour découvrir les méthodes de l’innovation de façon ludique, sans prérequis.",
     "les-dejeuxners-de-linnovation"),
    (JEU, "13:00", "19:00", "Pépinière numérique FabLabNab", "Quand la fabrication numérique booste l’innovation",
     "Visite", "", "INRAE, campus de la Grande Ferrade", "Villenave-d’Ornon", "FabLabNab, bâtiment Tripode, 71 avenue Édouard Bourlaux, 33140 Villenave-d’Ornon",
     "Chercheurs, enseignants-chercheurs, doctorants",
     "Visite du FabLabNab (impression 3D, découpe laser, prototypage rapide) et temps d’échange sur la science ouverte et l’innovation frugale.",
     "pepiniere-numerique-fablabnab-quand-la-fabrication-numerique-booste-linnovation"),
    (JEU, "14:00", "15:00", "La robotique au service de la fabrication additive métallique", "",
     "Visite", "", "Arts et Métiers, halle technologique", "Talence", ARTS_METIERS,
     "Enseignants-chercheurs, chercheurs, ingénieurs, doctorants, acteurs socio-économiques",
     "Présentation-visite d’une cellule robotisée de fabrication additive métallique (WAAM) : fonctionnement, usages industriels et défis.",
     "la-robotique-au-service-de-la-fabrication-additive-metallique"),
    (JEU, "17:30", "22:00", "Afterwork docteurs, doctorants et entreprises", "",
     "Afterwork", "", "Campus Pey-Berland, salle 1K", "Bordeaux", "35 place Pey Berland, 33000 Bordeaux",
     "Docteurs et doctorants",
     "Soirée de rencontre entre docteurs, doctorants et entreprises bordelaises : témoignages, échanges en petits groupes et réseautage.",
     "afterwork-docteurs-doctorants-et-entreprises"),
    # ------------------------------------------------------------- Vendredi 27
    (VEN, "09:00", "10:00", "CampusIABordeaux", "Entraînez vos étudiants sur la donnée",
     "Atelier", "", "BU Droit-Lettres, salle 205", "Pessac", "Salle 205, " + BU_DL,
     "Enseignants-chercheurs",
     "Présentation de la plateforme CampusIABordeaux et de son usage avec les étudiants : publication, partage et réutilisation des données.",
     "campusiabordeaux-entrainez-vos-etudiants-sur-la-donnee"),
    (VEN, "09:00", "11:00", "Des formations sur mesure pour faire évoluer ses pratiques", "",
     "Table ronde", "", "IMS, Salle de conférence", "Talence", "Bâtiment A31, " + IMS,
     "Accompagnants à l’innovation, enseignants-chercheurs",
     "Trois retours d’expérience sur des formations-accompagnements conçus sur mesure pour répondre à des besoins collectifs.",
     "des-formations-sur-mesure-pour-faire-evoluer-ses-pratiques"),
    (VEN, "11:00", "12:00", "Serious Game « No wrong door » du Pôle universitaire d’innovation Bordeaux", "",
     "Atelier", "", "IMS, Salle de conférence", "Talence", "Bâtiment A31, " + IMS,
     "Personnels administratifs",
     "Jeu en équipes de 4 à 5 personnes pour tester sa connaissance de l’écosystème d’innovation bordelais et savoir orienter les porteurs de projets.",
     "le-pole-universitaire-dinnovation-bordeaux-en-jeu-venez-relever-le-defi"),
    (VEN, "11:30", "12:30", "CampusIABordeaux", "Entraînez vos étudiants sur la donnée",
     "Atelier", "", "Campus Victoire, salle 32", "Bordeaux", "Salle 32, bâtiment B, 3 ter place de la Victoire, 33000 Bordeaux",
     "Enseignants-chercheurs",
     "Présentation de la plateforme CampusIABordeaux et de son usage avec les étudiants : publication, partage et réutilisation des données.",
     "campusiabordeaux-entrainez-vos-etudiants-sur-la-donnee"),
]


def construire(chemin):
    wb = Workbook()
    ws = wb.active
    ws.title = "Evenements"  # Grist nomme la table d'après la feuille

    police = "Arial"
    entete = PatternFill("solid", start_color="B9348B")
    for c, (nom, largeur) in enumerate(COLONNES, start=1):
        cell = ws.cell(row=1, column=c, value=nom)
        cell.font = Font(name=police, bold=True, color="FFFFFF")
        cell.fill = entete
        cell.alignment = Alignment(vertical="center")
        ws.column_dimensions[get_column_letter(c)].width = largeur

    lignes = lignes_triees()
    for r, (jour, debut, fin, titre, sous, type_, fort, lieu, ville, adresse, public, desc, slug) in enumerate(lignes, start=2):
        lien = PROGRAMME if slug is None else SITE + slug
        valeurs = [jour, debut, fin or None, titre, sous or None, type_, fort or None,
                   lieu, ville, adresse or None, public, cibles_de(jour, debut, titre) or None, desc or None, lien]
        for c, v in enumerate(valeurs, start=1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.font = Font(name=police)
            cell.alignment = Alignment(vertical="top", wrap_text=c in (4, 5, 8, 10, 11, 13))
        ws.cell(row=r, column=1).number_format = "yyyy-mm-dd"
        # heures en texte « 09:30 » : évite toute conversion de fuseau ou de format à l'import
        ws.cell(row=r, column=2).number_format = "@"
        ws.cell(row=r, column=3).number_format = "@"

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(chemin)
    return len(lignes)


def verifier_cibles():
    cles = {(e[0], e[1], e[3]) for e in EVENEMENTS}
    inconnues = [c for c in CIBLES if c not in cles]
    if inconnues:
        raise SystemExit(f"CIBLES : séance introuvable {inconnues}")


def lignes_triees():
    return sorted(EVENEMENTS, key=lambda e: (e[0], e[1], e[2] or "99:99", e[3]))


def ecrire_donnees_dev(chemin):
    """Mêmes lignes que l'Excel, au format d'une table Grist (Jour en secondes UTC), pour le faux Grist de dev/."""
    noms = [nom for nom, _ in COLONNES]
    lignes = []
    for i, e in enumerate(lignes_triees(), start=1):
        jour, debut, fin, titre, sous, type_, fort, lieu, ville, adresse, public, desc, slug = e
        cibles = [c for c in cibles_de(jour, debut, titre).split(", ") if c]
        valeurs = [timegm(jour.timetuple()), debut, fin, titre, sous, type_, fort, lieu, ville, adresse, public,
                   ["L", *cibles] if cibles else None,   # liste de choix : Grist la renvoie sous la forme ['L', 'a', 'b']
                   desc, PROGRAMME if slug is None else SITE + slug]
        lignes.append({"id": i, **dict(zip(noms, valeurs))})
    with open(chemin, "w", encoding="utf-8") as f:
        f.write("// Généré par data/build_xlsx.py — ne pas modifier à la main.\n")
        f.write("window.MOCK_EVENEMENTS = " + json.dumps(lignes, ensure_ascii=False, indent=1) + ";\n")


def ecrire_cibles_planning(chemin):
    """Les publics de chaque séance, lus par le widget pour remplir la colonne « Cibles » quand elle n'existe pas encore
    (ni vue du widget : il la crée alors, en propriétaire du document). Séance repérée par son jour, son début et son titre."""
    lignes = []
    for jour, debut, _fin, titre, *_ in lignes_triees():
        publics = [c for c in cibles_de(jour, debut, titre).split(", ") if c]
        if publics:
            lignes.append([jour.isoformat(), debut, titre, publics])
    with open(chemin, "w", encoding="utf-8") as f:
        f.write("// Généré par data/build_xlsx.py — ne pas modifier à la main.\n")
        f.write("// Publics de chaque séance d'après les plannings PDF : [jour, début, titre, [publics]].\n")
        f.write("window.CIBLES_PLANNING = [\n")
        f.write(",\n".join(" " + json.dumps(l, ensure_ascii=False) for l in lignes))
        f.write("\n];\n")


if __name__ == "__main__":
    verifier_cibles()
    sortie = sys.argv[1] if len(sys.argv) > 1 else "data/Evenements_Semaine_Innovation_2026.xlsx"
    print(f"{construire(sortie)} séances écrites dans {sortie}")
    ecrire_donnees_dev("dev/evenements.js")
    ecrire_cibles_planning("cibles-planning.js")
